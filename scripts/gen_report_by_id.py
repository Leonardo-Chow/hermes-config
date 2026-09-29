#!/usr/bin/env python3
"""按 video ID 生成上线监测报告 —— 标题直接从 YouTube API 取，杜绝手抄/截断误差。"""
import sys

sys.path.insert(0, "/Users/zhoulong/.hermes/scripts")
from docx import Document  # noqa: E402
import obsbot_yt_scan as s  # noqa: E402

OUT = "/Users/zhoulong/Downloads/2026-09-28-视频上线监测.docx"

DAYS = [
    ("2026/09/24（周四 UTC）", [
        ("Tiny 3 Lite", "NZL36keRycs"),
        ("Meet 2", "tbYmKs5vhhc"),
        ("Tiny 3 Lite", "2RW10c3N1So"),
        ("Meet Flip", "yMqZ2YxA2jA"),
        ("Tiny 3", "v2tUM6bOEC0"),
        ("Tail 2", "OeDMAAaMAeU"),
        ("Tiny 3", "blXAoxp9cx4"),
    ]),
    ("2026/09/25（周五 UTC）— OBSBOT Meet Flip 评测集中爆发（印尼区）", [
        ("Tail 2 & Tail Air", "1jg9nSNA2Zo"),
        ("Meet Flip", "0DamTtKbd9U"),
        ("Meet Flip", "DqVlI5oiJJA"),
        ("Meet Flip", "7ltWTV6XNZ4"),
        ("Meet Flip", "cuejdZre64s"),
        ("Meet Flip", "5uqXWYeWh3U"),
        ("Meet Flip", "LRIOeUGnLhc"),
        ("Meet Flip", "PU3ZCdhIHt0"),
        ("Meet Flip", "RTBe5aJu11I"),
        ("Meet Flip", "ddlciGiRHGI"),
        ("Tail 2", "hkJzQN8Jjc4"),
        ("Meet Flip", "i9gGRhyDosY"),
        ("Meet Flip", "eMVHJM8xNPU"),
        ("Meet Flip", "tF3f3nHChdQ"),
        ("Meet Flip", "MiUBowCy1gM"),
        ("Meet SE", "Rj9J130ipiM"),
    ]),
    ("2026/09/26（周六 UTC）", [
        ("Tiny 3 + VOX SE", "nHIT21Ji17c"),
        ("Talent 2", "Ul4jhBISOjQ"),
        ("Meet 2", "VOq3MN12Qho"),
        ("Tiny 2 Lite", "6M_rMAkzkK8"),
    ]),
    ("2026/09/27（周日 UTC）", [
        ("OBSBOT", "sjIqYx2bAeQ"),
        ("Talent 2", "JDeMoGGICJE"),
        ("Tiny 2 Lite", "U5HUqQHnNqQ"),
        ("OBSBOT Tiny 3 【赞助植入】", "Hp0QDOhe0p4"),
        ("Talent 2 【赞助植入】", "oXcHvVDPwt4"),
        ("Tiny 2 Lite 【赞助植入】", "N87W-59qKgc"),
    ]),
]

all_ids = [vid for _, items in DAYS for _, vid in items]
meta = {}
for k in range(0, len(all_ids), 50):
    d = s.api("videos", {"part": "snippet", "id": ",".join(all_ids[k:k + 50])})
    for it in d.get("items", []):
        meta[it["id"]] = (it["snippet"]["channelTitle"], it["snippet"]["title"])

missing = [v for v in all_ids if v not in meta]
if missing:
    sys.exit(f"API 未返回以下视频，终止: {missing}")

doc = Document()
doc.add_paragraph("2026/09/28 上线情况（补齐 9/24–9/27 共 4 天）")
doc.add_paragraph("")
for header, items in DAYS:
    doc.add_paragraph(header)
    for i, (prod, vid) in enumerate(items, 1):
        ch, title = meta[vid]
        doc.add_paragraph(f"{i}. {prod}-{ch}")
        doc.add_paragraph(f"https://www.youtube.com/watch?v={vid}")
        doc.add_paragraph(title)          # 原样标题，不截断不改写
        doc.add_paragraph("")
for sec in ["INS", "TT", "X"]:
    doc.add_paragraph(sec)
    doc.add_paragraph("暂无")
    doc.add_paragraph("")
doc.add_paragraph("搜索关键词")
doc.add_paragraph("OBSBOT / OBSBOT Tiny 3 / OBSBOT Tiny 3 Lite / OBSBOT Tiny 2 / OBSBOT Tiny 2 Lite / "
                  "OBSBOT Tail 2 / OBSBOT Tail Air / OBSBOT Meet Flip / OBSBOT Meet SE / OBSBOT Meet 2 / "
                  "OBSBOT Meet / OBSBOT Talent / OBSBOT Vox / OBSBOT review / OBSBOT unboxing / OBSBOT webcam")
doc.save(OUT)
print(f"已重新生成: {OUT}（{len(all_ids)} 条，标题全部取自 API）")
